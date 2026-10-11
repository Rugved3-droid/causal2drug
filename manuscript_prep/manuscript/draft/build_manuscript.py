#!/usr/bin/env python
"""Build manuscript.md / manuscript_FINAL.docx and supplement.docx, and count words.

    python manuscript_prep/manuscript/draft/fetch_refs.py        # once; resolves DOIs
    python manuscript_prep/manuscript/draft/build_manuscript.py

manuscript_src.md carries citations as {{key,key}}; they are numbered in order of first
citation. Every reference is rendered from Crossref metadata (refs_crossref.json), so a DOI
that did not resolve cannot be cited. Word count = abstract + main text + figure legends,
as the journal defines it; headings are counted, citations and tables are not.
"""
import json
import os
import re
import sys

import pypandoc

sys.stdout.reconfigure(encoding="utf-8")
HERE = os.path.dirname(os.path.abspath(__file__))
MS = os.path.dirname(HERE)
MP = os.path.dirname(MS)
KEYS = {
    "minikel": "10.1038/s41586-024-07316-0", "nelson": "10.1038/ng.3314",
    "king": "10.1371/journal.pgen.1008489", "ochoa": "10.1038/d41573-022-00120-3",
    "rusina": "10.1038/d41573-023-00158-x", "schmidt": "10.1038/s41467-020-16969-0",
    "gill21": "10.12688/wellcomeopenres.16544.2", "gill24": "10.1186/s12916-024-03700-9",
    "zheng": "10.1038/s41588-020-0682-6", "vosa": "10.1038/s41588-021-00913-z",
    "storm": "10.1038/s41467-021-26280-1", "sadler": "10.1016/j.xgen.2023.100341",
    "gaziano": "10.1038/s41591-021-01310-z", "song": "10.1186/s12885-024-12449-6",
    "zhang": "10.1016/j.ebiom.2023.104859", "liang": "10.1038/s41598-025-94761-0",
    "liu": "10.1186/s12931-024-02848-5", "wang": "10.3389/fcimb.2024.1386506",
    "gtex": "10.1126/science.aaz1776", "sun": "10.1038/s41586-023-06592-6",
    "ferkingstad": "10.1038/s41588-021-00978-w", "pietzner": "10.1126/science.abj1541",
    "finan": "10.1126/scitranslmed.aag1166", "mostafavi": "10.1038/s41588-023-01529-1",
    "umans": "10.1016/j.tig.2020.08.009", "wise": "10.1016/j.ajhg.2013.03.017",
    "moix": "10.1186/s13073-026-01689-9", "ferolito": "10.1016/j.xhgg.2025.100556",
    "strobe": "10.1001/jama.2021.18236", "chembl": "10.1093/nar/gkad1004",
    "hpa": "10.1126/science.1260419", "gordillo": "10.1038/s43856-024-00625-5",
    "coloc": "10.1371/journal.pgen.1004383", "hgnc": "10.1093/nar/gkac888",
    "rosen": "10.1101/2025.08.05.668745", "burgess": "10.12688/wellcomeopenres.15555.3", "burgess16": "10.1002/gepi.21998",
    "ferreira": "10.1371/journal.pgen.1003444",
    "cohen": "10.1056/NEJMoa054013", "stitziel": "10.1056/NEJMoa1405386",
    "tghdl": "10.1056/NEJMoa1307095", "jorgensen": "10.1056/NEJMoa1308027",
    "dewey": "10.1056/NEJMoa1612790", "ference": "10.1056/NEJMoa1604304",
    "il6rmr": "10.1016/S0140-6736(12)60110-X", "clarke": "10.1056/NEJMoa0902604",
    "schmitt": "10.1126/science.1081578", "towbin": "10.1161/01.CIR.87.6.1854"}
# Corrections to Crossref metadata, each checked against PubMed on 2026-10-11.
AUTHOR_FIX = {
    "10.1038/d41573-022-00120-3": ["Ochoa D", "Karim M", "Ghoussaini M", "Hulcoop DG",
                                   "McDonagh EM", "Dunham I"],
    "10.1016/S0140-6736(12)60110-X": [
        "Interleukin-6 Receptor Mendelian Randomisation Analysis (IL6R MR) Consortium"],
    "10.1056/NEJMoa1405386": ["Myocardial Infarction Genetics Consortium Investigators"],
    "10.1056/NEJMoa1307095": ["TG and HDL Working Group of the Exome Sequencing Project, "
                              "National Heart, Lung, and Blood Institute"],
}
# The MR guidelines are version 3 (2023) of a paper first published in 2019.
NOTE = {"10.12688/wellcomeopenres.15555.3": ("2019", " [version 3, 2023]")}
JABBR = {"The Lancet": "Lancet", "The American Journal of Human Genetics": "Am J Hum Genet",
         "Human Genetics and Genomics Advances": "HGG Adv", "Trends in Genetics": "Trends Genet",
         "Nucleic Acids Research": "Nucleic Acids Res", "Cell Genomics": "Cell Genom",
         "Front. Cell. Infect. Microbiol.": "Front Cell Infect Microbiol",
         "Sci. Transl. Med.": "Sci Transl Med", "eBioMedicine": "EBioMedicine"}


def to_docx(md_path, docx_path):
    """Convert with pandoc. If the target is open in Word (Windows locks it), write beside it
    under a new name instead of failing, and say so."""
    args = ["--from=markdown+superscript", "--resource-path=" + MP]
    try:
        pypandoc.convert_file(md_path, "docx", outputfile=docx_path, extra_args=args)
        return docx_path
    except RuntimeError as e:
        if "permission denied" not in str(e).lower():
            raise
        alt = docx_path.replace(".docx", "_v2.docx")
        pypandoc.convert_file(md_path, "docx", outputfile=alt, extra_args=args)
        print("NOTE: %s is open in another program; wrote %s instead"
              % (os.path.basename(docx_path), os.path.basename(alt)))
        return alt


REFS = json.load(open(os.path.join(HERE, "refs_crossref.json"), encoding="utf-8"))


def fmt_ref(doi):
    r = REFS[doi]                       # KeyError = unverified reference; never cited
    au = [a.replace("​", "").strip() for a in r["authors"] if a.strip()]
    au = AUTHOR_FIX.get(doi, au)
    au = [a for a in au if a != "The GTEx Consortium"] if doi != "10.1126/science.aaz1776" \
        else ["GTEx Consortium"]
    names = ", ".join(au) if len(au) <= 6 else ", ".join(au[:3]) + ", et al"
    j = JABBR.get(r["journal"], r["journal"])
    if doi.startswith("10.1101/"):
        return "%s. %s. *bioRxiv*. Preprint posted %s. doi: %s" % (
            names, r["title"].rstrip("."), r["year"], doi)
    pg = r["page"]
    if pg and "-" in pg and pg.split("-")[0] == pg.split("-")[1]:
        pg = pg.split("-")[0]                      # single-page items: 551, not 551-551
    vol = "%s:%s" % (r["volume"], pg) if pg else str(r["volume"])
    t = r["title"].rstrip(".")
    if doi in NOTE:
        r = dict(r, year=NOTE[doi][0])
        vol = vol + NOTE[doi][1]
    return "%s. %s%s *%s*. %s;%s. doi: %s" % (names, t, "" if t.endswith("?") else ".", j,
                                              r["year"], vol, doi)


def compress(ns):
    ns = sorted(set(ns))
    out, i = [], 0
    while i < len(ns):
        j = i
        while j + 1 < len(ns) and ns[j + 1] == ns[j] + 1:
            j += 1
        out.append(str(ns[i]) if j == i else "%d,%d" % (ns[i], ns[j]) if j == i + 1
                   else "%d–%d" % (ns[i], ns[j]))
        i = j + 1
    return ",".join(out)


def cite_and_number(text, order=None):
    order = [] if order is None else order

    def sub(m):
        ks = [k.strip() for k in m.group(1).split(",")]
        for k in ks:
            if KEYS[k] not in order:
                order.append(KEYS[k])
        return "^%s^" % compress([order.index(KEYS[k]) + 1 for k in ks])
    return re.sub(r"\{\{([a-z0-9, ]+)\}\}", sub, text), order


def words(s):
    s = re.sub(r"<!--.*?-->", " ", s, flags=re.S)
    s = re.sub(r"\{\{[a-z0-9, ]+\}\}", " ", s)
    s = re.sub(r"\^[0-9,–]+\^", " ", s)
    s = re.sub(r"[*#_`]", "", s)
    return len([w for w in s.split() if re.search(r"[A-Za-z0-9]", w)])


def table_md(path, drop_title=False):
    t = open(path, encoding="utf-8").read().strip()
    return t


src = open(os.path.join(HERE, "manuscript_src.md"), encoding="utf-8").read()
counts = {}
for name in ("abstract", "main", "legends"):
    m = re.search(r"<!--COUNT:%s-->(.*?)<!--END-->" % name, src, flags=re.S)
    counts[name] = words(m.group(1))
abs_body = re.search(r"<!--COUNT:abstract-->(.*?)\*\*Key Words", src, flags=re.S).group(1)
counts["abstract_body_only"] = words(abs_body.replace("## Abstract", ""))
counts["abstract"] = counts["abstract_body_only"]          # key words are not counted
counts["total"] = counts["abstract"] + counts["main"] + counts["legends"]
sec = re.search(r"<!--COUNT:main-->(.*?)<!--END-->", src, flags=re.S).group(1)
bysec = [(h.strip(), words(b)) for h, b in re.findall(r"\n## ([^\n]+)\n(.*?)(?=\n## |\Z)", sec,
                                                      flags=re.S)]

body, order = cite_and_number(src)
reflist = "\n".join("%d. %s" % (i + 1, fmt_ref(d)) for i, d in enumerate(order))
tabs = "\n\n".join(table_md(p) for p in (
    os.path.join(MP, "tables", "table1_flow_and_cohort.md"),
    os.path.join(MS, "display", "table2_named_targets.md"),
    os.path.join(MS, "display", "table3_recovery_tiers.md")))
body = body.replace("[[REFERENCES]]", reflist).replace("[[TABLES]]", tabs)
body = body.replace("[[WORDCOUNT]]", "{:,}".format(counts["total"]))
body = re.sub(r"<!--.*?-->\n?", "", body)
out_md = os.path.join(MS, "manuscript.md")
open(out_md, "w", encoding="utf-8", newline="\n").write(body)
DOCX = to_docx(out_md, os.path.join(MS, "manuscript_FINAL.docx"))


def landscape_table2(path):
    """Put Table 2 (12 columns) on its own landscape section and set its column widths."""
    import copy
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt
    d = Document(path)
    body = d.element.body
    kids = list(body.iterchildren())

    def text(el):
        return "".join(t.text or "" for t in el.iter(qn("w:t")))
    i2 = next(i for i, k in enumerate(kids) if k.tag == qn("w:p") and text(k).startswith("Table 2."))
    i3 = next(i for i, k in enumerate(kids) if k.tag == qn("w:p") and text(k).startswith("Table 3."))
    final = body.find(qn("w:sectPr"))

    def sect(landscape):
        sp = copy.deepcopy(final)
        pg = sp.find(qn("w:pgSz"))
        if pg is None:
            pg = OxmlElement("w:pgSz")
            sp.insert(0, pg)
        w, h = ("15840", "12240") if landscape else ("12240", "15840")
        pg.set(qn("w:w"), w)
        pg.set(qn("w:h"), h)
        if landscape:
            pg.set(qn("w:orient"), "landscape")
        elif pg.get(qn("w:orient")) is not None:
            del pg.attrib[qn("w:orient")]
        return sp

    def end_section(par, landscape):
        ppr = par.find(qn("w:pPr"))
        if ppr is None:
            ppr = OxmlElement("w:pPr")
            par.insert(0, ppr)
        ppr.append(sect(landscape))
    prev2 = next(k for k in reversed(kids[:i2]) if k.tag == qn("w:p"))
    prev3 = next(k for k in reversed(kids[:i3]) if k.tag == qn("w:p"))
    end_section(prev2, False)          # everything before Table 2: portrait
    end_section(prev3, True)           # Table 2 with its footnotes: landscape
    tbl_el = next(k for k in kids[i2:i3] if k.tag == qn("w:tbl"))
    tbl = next(t for t in d.tables if t._tbl is tbl_el)
    ncol = len(tbl.columns)
    rel = [0.62, 1.15, 1.2, 1.05, 0.8, 0.8, 0.42, 1.55, 0.45, 0.62, 0.62, 0.62][:ncol]
    rel += [0.6] * (ncol - len(rel))
    scale = 9.0 / sum(rel)             # 11 in wide page, 1 in margins
    tbl.autofit = False
    for row in tbl.rows:
        for j, cell in enumerate(row.cells):
            cell.width = Inches(rel[j] * scale)
            for par in cell.paragraphs:
                for run in par.runs:
                    run.font.size = Pt(6.5)
    d.save(path)
    return ncol


try:
    print("Table 2 set on a landscape section; columns: %d" % landscape_table2(DOCX))
except Exception as exc:                                         # noqa: BLE001
    print("NOTE: landscape step failed: %r" % (exc,))
print("manuscript written to %s" % os.path.basename(DOCX))

n_display = 6
wc = ["# WORD_COUNT", "",
      "Counted by `python manuscript_prep/manuscript/draft/build_manuscript.py` from "
      "`draft/manuscript_src.md`. Citation markers, tables, title page, key words, "
      "abbreviations, references and end statements are not counted; section headings are.",
      "", "| part | words | limit |", "|---|---|---|",
      "| Structured abstract | %d | 250 |" % counts["abstract"],
      "| Main text (Introduction to Conclusions, with headings) | %d | |" % counts["main"],
      "| Figure legends | %d | |" % counts["legends"],
      "| **Total (abstract + main text + legends)** | **%d** | **5,000** |" % counts["total"],
      "| References | %d | 50 |" % len(order),
      "| Display items (3 tables, 3 figures) | %d | 8 |" % n_display, "",
      "## Main text by section", "", "| section | words |", "|---|---|"]
wc += ["| %s | %d |" % kv for kv in bysec]
wc += ["", "Status: abstract %s; total %s; references %s; display items %s." % (
    "within limit" if counts["abstract"] <= 250 else "OVER LIMIT",
    "within limit" if counts["total"] <= 5000 else "OVER LIMIT",
    "within limit" if len(order) <= 50 else "OVER LIMIT",
    "within limit" if n_display <= 8 else "OVER LIMIT")]
open(os.path.join(MS, "WORD_COUNT.md"), "w", encoding="utf-8", newline="\n").write(
    "\n".join(wc) + "\n")
print("\n".join(wc[4:12]))
print("cited references:", len(order), "| not cited:",
      [k for k, d in KEYS.items() if d not in order])

# ---------------------------------------------------------------- supplement
sp = os.path.join(HERE, "supplement_src.md")
if os.path.exists(sp):
    s = open(sp, encoding="utf-8").read()

    def inc(m):
        p = os.path.join(MP, m.group(1))
        t = open(p, encoding="utf-8").read().strip().split("\n")
        lim = int(m.group(2)) if m.group(2) else None
        if lim:
            head = [i for i, l in enumerate(t) if l.startswith("|")]
            rows = head[2:]
            if len(rows) > lim:
                cut = rows[lim]
                last = head[-1]
                t = t[:cut] + ["", "*First %d of %d rows shown; the full table is in the "
                               "supplemental workbook and CSV file.*" % (lim, len(rows)), ""] \
                    + t[last + 1:]
        return "\n".join(t)
    s = re.sub(r"\[\[INCLUDE:([^\]|]+)(?:\|(\d+))?\]\]", inc, s)
    s, order2 = cite_and_number(s, list(order))
    extra = order2[len(order):]
    s = s.replace("[[SUPPREFS]]", "\n".join(
        "%d. %s" % (len(order) + i + 1, fmt_ref(d)) for i, d in enumerate(extra))
        or "None beyond the main reference list.")
    out_s = os.path.join(MS, "supplement.md")
    open(out_s, "w", encoding="utf-8", newline="\n").write(s)
    to_docx(out_s, os.path.join(MS, "supplement.docx"))
    print("supplement built; extra references:", len(extra))
