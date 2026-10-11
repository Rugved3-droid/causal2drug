#!/usr/bin/env python
"""Print sheet names, dimensions and the first rows of a workbook (headers only).
    python -I inspect_xlsx.py <workbook> [nrows] [sheet-substring]
"""
import sys
import openpyxl
sys.stdout.reconfigure(encoding="utf-8")
wb = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)
n = int(sys.argv[2]) if len(sys.argv) > 2 else 6
sub = sys.argv[3] if len(sys.argv) > 3 else ""
for ws in wb.worksheets:
    if sub and sub.lower() not in ws.title.lower():
        continue
    print("=== %s  rows=%s cols=%s" % (ws.title, ws.max_row, ws.max_column))
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i >= n:
            break
        print("   ", [("" if c is None else str(c)[:60]) for c in row][:40])
